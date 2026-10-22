"""Collecting result rows and writing them to results.csv and the course workbook.

results.csv is the machine-readable record kept in the repository. It has a
"Sheet" column followed by the template columns of every sheet that has rows,
in template order. The workbook is a filled copy of the course template.
"""

from __future__ import annotations

import datetime as dt
import math
from collections.abc import Mapping
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from recsys.config import TeamConfig
from recsys.reporting.schema import FIRST_DATA_ROW, HEADER_ROW, INSTRUCTION_CELLS, INSTRUCTIONS_SHEET, METRIC_COLUMNS, SHEETS
from recsys.utils import save_table


class ResultsTable:
    def __init__(self) -> None:
        self._rows: list[tuple[str, dict[str, Any]]] = []

    def add(self, sheet: str, row: Mapping[str, Any]) -> None:
        if sheet not in SHEETS:
            raise KeyError(f"Unknown results sheet {sheet!r}")
        unknown = sorted(set(row) - set(SHEETS[sheet]))
        if unknown:
            raise KeyError(f"Columns {unknown} are not in the template sheet {sheet!r}")
        self._rows.append((sheet, {col: row.get(col) for col in SHEETS[sheet]}))

    def __len__(self) -> int:
        return len(self._rows)

    def rows(self, sheet: str) -> list[dict[str, Any]]:
        return [row for name, row in self._rows if name == sheet]

    def frame(self) -> pd.DataFrame:
        columns: list[str] = []
        for sheet, cols in SHEETS.items():
            if self.rows(sheet):
                columns += [c for c in cols if c not in columns]
        records = [{"Sheet": sheet, **row} for sheet, row in self._rows]
        return pd.DataFrame.from_records(records, columns=["Sheet", *columns])

    def write_csv(self, path: Path) -> None:
        save_table(self.frame(), path, float_format="%.6f")

    def write_workbook(self, template: Path, path: Path, team: TeamConfig) -> None:
        # keep_vba keeps the macro-enabled package type, so the .xlsm opens in Excel.
        wb = load_workbook(template, keep_vba=True)
        _fill_instructions(wb[INSTRUCTIONS_SHEET], team)
        for sheet, columns in SHEETS.items():
            ws = wb[sheet]
            header = tuple(ws.cell(HEADER_ROW, c + 1).value for c in range(len(columns)))
            if header != columns:
                raise ValueError(f"Template sheet {sheet!r} has header {header}, expected {columns}")
            for r, row in enumerate(self.rows(sheet), start=FIRST_DATA_ROW):
                for c, col in enumerate(columns, start=1):
                    value = _cell_value(row[col])
                    cell = ws.cell(r, c, value)
                    if col in METRIC_COLUMNS and isinstance(value, float):
                        cell.number_format = "0.0000"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)


def _cell_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if hasattr(value, "item"):  # numpy scalar
        return value.item()
    return value


def _fill_instructions(ws, team: TeamConfig) -> None:
    ws[INSTRUCTION_CELLS["team"]] = team.name
    ws[INSTRUCTION_CELLS["members"]] = ", ".join(team.members)
    if team.repository:
        ws[INSTRUCTION_CELLS["repository"]] = team.repository
    if team.part3_track:
        ws[INSTRUCTION_CELLS["part3_track"]] = team.part3_track
    ws[INSTRUCTION_CELLS["last_updated"]] = dt.date.today().isoformat()
