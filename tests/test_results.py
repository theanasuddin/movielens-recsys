from __future__ import annotations

import pytest
from openpyxl import load_workbook

from recsys.config import TeamConfig
from recsys.reporting import SHEETS, ResultsTable
from recsys.reporting.schema import FIRST_DATA_ROW, HEADER_ROW
from recsys.utils import REPO_ROOT

TEMPLATE = REPO_ROOT / "templates" / "Recommender_Systems_Results_Template.xlsm"


def _row(**overrides):
    row = {
        "Model": "MostPopular", "Variant / settings": "all-time, all ratings", "Precision@10": 0.1,
        "Recall@10": 0.05, "NDCG@10": 0.12, "Catalogue coverage": 0.01, "Seed": 42, "Notes": "test",
    }
    row.update(overrides)
    return row


def test_template_headers_match_schema():
    wb = load_workbook(TEMPLATE, keep_vba=True)
    for sheet, columns in SHEETS.items():
        header = tuple(wb[sheet].cell(HEADER_ROW, c + 1).value for c in range(len(columns)))
        assert header == columns, sheet


def test_unknown_sheet_or_column_is_rejected():
    table = ResultsTable()
    with pytest.raises(KeyError):
        table.add("Part 9", _row())
    with pytest.raises(KeyError):
        table.add("Part 1 - Baselines", _row(NDCG=0.1))


def test_csv_has_sheet_column_and_template_columns(tmp_path):
    table = ResultsTable()
    table.add("Part 1 - Baselines", _row())
    path = tmp_path / "results.csv"
    table.write_csv(path)
    header = path.read_text(encoding="utf-8").splitlines()[0].split(",")
    assert header == ["Sheet", *SHEETS["Part 1 - Baselines"]]


def test_workbook_rows_land_under_the_header(tmp_path):
    table = ResultsTable()
    table.add("Part 1 - Baselines", _row())
    table.add("Part 1 - Baselines", _row(Model="Random"))
    out = tmp_path / "results.xlsm"
    table.write_workbook(TEMPLATE, out, TeamConfig(name="Group 2", members=("A", "B")))

    wb = load_workbook(out, keep_vba=True)
    ws = wb["Part 1 - Baselines"]
    assert ws.cell(FIRST_DATA_ROW, 1).value == "MostPopular"
    assert ws.cell(FIRST_DATA_ROW + 1, 1).value == "Random"
    assert ws.cell(FIRST_DATA_ROW, 7).value == pytest.approx(0.12)
    assert wb["Instructions"]["B3"].value == "Group 2"
    assert wb["Instructions"]["B4"].value == "A, B"
